"""Bounded CSV/XLSX reader. Uses the existing openpyxl dependency, no customer import changes."""
import csv
import io
import json
import zipfile
from datetime import date,datetime,time
from fastapi import HTTPException
from openpyxl import load_workbook
from ..schemas.order_sources import MAX_BYTES,MAX_ROWS,ExternalOrder
from .order_grouping import normalize

FIELDS=list(ExternalOrder.model_fields)
ALIASES={'external_id':['idordine','id','nordine','numeroordine'], 'number':['numeroordine','nordine','numero'],
 'external_customer_id':['codcli','codicecliente'], 'recipient_name':['ragionesociale','cliente','destinatario'],
 'delivery_address':['indsped','indirizzo','indirizzoconsegna'], 'requested_date':['datacons','dataconsegna'],
 'weight_kg':['peso','pesokg'],'packages':['ncolli','colli'],'tail_lift':['pedana','sponda'],
 'pallet_truck':['transpallet'],'ztl':['ztl'],'notes':['note'],'time_from':['dalle'],'time_to':['alle'],
 'pallets':['pallet'],'volume_m3':['volume'],'requirements':['requisiti']}


def suggestions(headers):
    def key(value): return normalize(value).replace(' ','').replace('_','')
    result={}
    for target in FIELDS:
        candidates=[h for h in headers if key(h) in {key(target),*(key(a) for a in ALIASES.get(target,[]))}]
        if len(candidates)==1: result[target]=candidates[0]
    if 'external_id' not in result and 'number' in result: result['external_id']=result['number']
    return result


def validate_rows(headers, rows):
    if not headers or len(headers)>50 or any(not h or len(h)>160 for h in headers) or len(set(headers))!=len(headers):
        raise HTTPException(422,'Usa fino a 50 intestazioni univoche e non vuote (massimo 160 caratteri).')
    if not rows or len(rows)>MAX_ROWS: raise HTTPException(422,f'Il file deve contenere da 1 a {MAX_ROWS} ordini. Suddividi i file più grandi.')
    for row in rows:
        if len(row)>50 or any(len(str(v))>5000 for v in row.values()): raise HTTPException(422,'Troppi campi o una cella supera 5000 caratteri.')
    return headers,rows


def read_file(content,filename):
    if len(content)>MAX_BYTES: raise HTTPException(413,'Il file supera il limite tecnico di 2 MB.')
    try:
        if filename.lower().endswith('.csv'):
            try: text=content.decode('utf-8-sig')
            except UnicodeDecodeError: text=content.decode('cp1252')
            try: dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            except csv.Error: dialect=csv.excel
            reader=csv.reader(io.StringIO(text),dialect)
            headers=[h.strip() for h in next(reader)]
            rows=[]
            for values in reader:
                if not any(v.strip() for v in values): continue
                if len(values)!=len(headers): raise ValueError('columns')
                rows.append(dict(zip(headers,values)))
                if len(rows)>MAX_ROWS: raise ValueError('rows')
        elif filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                files=archive.infolist()
                if len(files)>1000 or sum(f.file_size for f in files)>20*1024*1024 or any('externallinks/' in f.filename.lower() or 'vbaproject' in f.filename.lower() for f in files):
                    raise ValueError('unsafe archive')
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                if len(book.worksheets)!=1: raise ValueError('Use one sheet')
                sheet=book.worksheets[0]
                if (sheet.max_row or 0)>MAX_ROWS+1 or (sheet.max_column or 0)>50: raise ValueError('dimensions')
                data=[]
                for cells in sheet.iter_rows():
                    values=[]
                    for cell in cells:
                        if cell.data_type in ('f','e'): raise ValueError('formula/error')
                        value=cell.value
                        values.append(value.isoformat() if isinstance(value,(datetime,date,time)) else value)
                    if any(v is not None for v in values): data.append(values)
                    if len(data)>MAX_ROWS+1: raise ValueError('rows')
                headers=[str(v or '').strip() for v in data[0]]
                rows=[dict(zip(headers,r)) for r in data[1:]]
            finally: book.close()
        else: raise HTTPException(422,'Usa un file CSV o XLSX con un solo foglio, senza formule.')
        return validate_rows(headers,rows)
    except HTTPException: raise
    except Exception:
        raise HTTPException(422,'File non leggibile o fuori limite: verifica intestazioni, colonne, massimo 1000 righe, un solo foglio XLSX e assenza di formule/link esterni.')


def mapped(row,mapping):
    data={key:row.get(column) for key,column in mapping.items() if key in FIELDS}
    for key,value in list(data.items()):
        if isinstance(value,str): value=value.strip()
        if value=='': value=None
        if key in ('external_id','number','external_customer_id') and type(value) in (int,float):
            if not float(value).is_integer() or abs(value)>=10**15: raise ValueError('Usa codici in formato testo.')
            value=str(int(value))
        if value is not None and key in ('tail_lift','pallet_truck','ztl') and isinstance(value,str):
            true={'true','si','sì','yes','1'};false={'false','no','0'}
            if value.casefold() not in true|false: raise ValueError(f'{key}: usa Sì/No o True/False.')
            value=value.casefold() in true
        if value is not None and key in ('weight_kg','pallets','volume_m3','packages') and isinstance(value,str):
            value=value.replace(',','.')
        if value is not None and key=='requested_date' and isinstance(value,str):
            if '/' in value: value=datetime.strptime(value,'%d/%m/%Y').date().isoformat()
            elif 'T' in value: value=datetime.fromisoformat(value).date().isoformat()
        if value is not None and key=='items' and isinstance(value,str): value=json.loads(value)
        data[key]=value
    if data.get('items') is None: data['items']=[]
    return ExternalOrder.model_validate(data).model_dump(mode='json')
